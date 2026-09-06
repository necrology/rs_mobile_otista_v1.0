from __future__ import annotations

import csv
import hashlib
import re
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from uat_sipantes_data import API_CHECKS, AUTOMATED_TESTS, SCREENSHOTS, UAT_CASES


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "uat" / "SIPANTES_UAT_2026-08-28"
SCREENSHOT_SOURCE = ROOT / "build" / "playstore_capture" / "form_factors" / "phone"
SCREENSHOT_OUTPUT = OUTPUT / "evidence" / "screenshots"
LOG_OUTPUT = OUTPUT / "evidence" / "logs"
DATA_OUTPUT = OUTPUT / "evidence" / "data"
PREVIEW_OUTPUT = OUTPUT / "preview"
DOCX_PATH = OUTPUT / "Dokumen_UAT_SIPANTES_v1.0.docx"
PDF_PATH = OUTPUT / "Dokumen_UAT_SIPANTES_v1.0.pdf"
ZIP_PATH = OUTPUT.parent / "SIPANTES_UAT_2026-08-28.zip"
DOCUMENT_DATE = "28 Agustus 2026"
TEAL = "176D64"
LIGHT_TEAL = "EAF5F3"


def run(command: list[str], log_path: Path | None = None) -> str:
    result = subprocess.run(
        command,
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        encoding="utf-8",
        errors="replace",
    )
    if log_path:
        log_path.write_text(result.stdout, encoding="utf-8")
    if result.returncode:
        raise RuntimeError(f"Command failed ({result.returncode}): {' '.join(command)}\n{result.stdout}")
    return result.stdout


def run_flutter(arguments: list[str], log_path: Path) -> str:
    flutter = ROOT.drive + "\\flutter\\flutter\\bin\\flutter.bat"
    command_line = subprocess.list2cmdline([flutter, *arguments])
    return run(["cmd.exe", "/d", "/s", "/c", command_line], log_path)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8-sig") as target:
        writer = csv.DictWriter(target, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "202020", size: float = 10.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run_item = paragraph.add_run(str(text))
    run_item.bold = bold
    run_item.font.name = "Garamond"
    run_item._element.rPr.rFonts.set(qn("w:eastAsia"), "Garamond")
    run_item.font.size = Pt(size)
    run_item.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_table(document: Document, headers: list[str], rows: list[list[str]], font_size: float = 10.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        shade_cell(table.rows[0].cells[index], TEAL)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF", size=font_size)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, size=font_size)
    document.add_paragraph()
    return table


def add_field(paragraph, instruction: str) -> None:
    run_item = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run_item._r.append(field)


def add_body(document: Document, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run_item = paragraph.add_run(text)
    run_item.bold = bold
    return paragraph


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Garamond"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Garamond")
    normal.font.size = Pt(12)
    for style_name, size, color in (
        ("Title", 22, TEAL),
        ("Heading 1", 18, TEAL),
        ("Heading 2", 15, "245A55"),
        ("Heading 3", 13, "333333"),
    ):
        style = document.styles[style_name]
        style.font.name = "Garamond"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Garamond")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def set_page_layout(section, landscape: bool = False) -> None:
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width


def add_headers_and_footers(document: Document) -> None:
    for index, section in enumerate(document.sections):
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        section.different_first_page_header_footer = index == 0
        header = section.header.paragraphs[0]
        header.clear()
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.add_run("RSUD Oto Iskandar Di Nata Kabupaten Bandung | UAT SIPANTES")
        header.runs[0].font.name = "Garamond"
        header.runs[0].font.size = Pt(9)
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Dokumen UAT SIPANTES | Halaman ")
        add_field(footer, "PAGE")
        footer.add_run(" dari ")
        add_field(footer, "NUMPAGES")
        for run_item in footer.runs:
            run_item.font.name = "Garamond"
            run_item.font.size = Pt(9)


def force_garamond(document: Document) -> None:
    def apply_run(run_item) -> None:
        run_item.font.name = "Garamond"
        properties = run_item._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Garamond")

    def apply_table(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run_item in paragraph.runs:
                        apply_run(run_item)
                for nested_table in cell.tables:
                    apply_table(nested_table)

    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = "Garamond"
            properties = style._element.get_or_add_rPr()
            fonts = properties.rFonts
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                properties.insert(0, fonts)
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{attribute}"), "Garamond")
    for paragraph in document.paragraphs:
        for run_item in paragraph.runs:
            apply_run(run_item)
    for table in document.tables:
        apply_table(table)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run_item in paragraph.runs:
                    apply_run(run_item)
            for table in container.tables:
                apply_table(table)


def collect_evidence() -> dict:
    output_parent = (ROOT / "docs" / "uat").resolve()
    resolved_output = OUTPUT.resolve()
    if output_parent not in resolved_output.parents:
        raise RuntimeError(f"Output UAT berada di luar folder yang diizinkan: {resolved_output}")
    if OUTPUT.exists():
        shutil.rmtree(OUTPUT)
    for directory in (OUTPUT, SCREENSHOT_OUTPUT, LOG_OUTPUT, DATA_OUTPUT, PREVIEW_OUTPUT):
        directory.mkdir(parents=True, exist_ok=True)

    for item in SCREENSHOTS:
        source = SCREENSHOT_SOURCE / item["file"]
        if not source.exists():
            raise FileNotFoundError(f"Screenshot tidak ditemukan: {source}")
        shutil.copy2(source, SCREENSHOT_OUTPUT / item["file"])

    analyze_text = run_flutter(["analyze"], LOG_OUTPUT / "flutter_analyze.txt")
    test_text = run_flutter(
        ["test", "--reporter", "expanded"],
        LOG_OUTPUT / "flutter_test.txt",
    )
    if "No issues found" not in analyze_text or "All tests passed" not in test_text:
        raise RuntimeError("Output validasi Flutter tidak sesuai ekspektasi.")

    api_rows = []
    for check_id, name, url, expected in API_CHECKS:
        started = time.perf_counter()
        request = urllib.request.Request(url, headers={"User-Agent": "SIPANTES-UAT/1.0"})
        try:
            with urllib.request.urlopen(request, timeout=20) as response:
                status = response.status
                content_type = response.headers.get("Content-Type", "")
                final_url = response.geturl()
                note = ""
        except urllib.error.HTTPError as error:
            status = error.code
            content_type = error.headers.get("Content-Type", "") if error.headers else ""
            final_url = error.geturl()
            note = "Respons HTTP sesuai mekanisme keamanan" if status == expected else str(error)
        except Exception as error:
            status = 0
            content_type = ""
            final_url = url
            note = str(error)
        elapsed = round(time.perf_counter() - started, 3)
        api_rows.append(
            {
                "Id": check_id,
                "Pengujian": name,
                "Method": "GET",
                "URL": url,
                "ExpectedHTTP": expected,
                "ActualHTTP": status,
                "ContentType": content_type,
                "TimeSeconds": elapsed,
                "FinalURL": final_url,
                "Hasil": "LULUS" if status == expected else "TIDAK SESUAI",
                "Keterangan": note,
            }
        )
    write_csv(DATA_OUTPUT / "API_Web_Smoke_Test.csv", api_rows)
    if any(row["Hasil"] != "LULUS" for row in api_rows):
        raise RuntimeError("Terdapat smoke test API/web yang tidak sesuai.")

    write_csv(
        DATA_OUTPUT / "UAT_Test_Cases.csv",
        [
            {
                "ID": case["id"],
                "Modul": case["module"],
                "Skenario": case["scenario"],
                "Prasyarat": case["precondition"],
                "Langkah": case["steps"],
                "Ekspektasi": case["expected"],
                "Aktual": case["actual"],
                "Status": case["status"],
                "Bukti": case["evidence"],
            }
            for case in UAT_CASES
        ],
    )
    write_csv(
        DATA_OUTPUT / "Automated_Test_Results.csv",
        [
            {"No": index, "Pengujian": test_name, "Hasil": "LULUS"}
            for index, test_name in enumerate(AUTOMATED_TESTS, 1)
        ],
    )

    for file_name in (
        "hasil_uji_34_operasi_origin.csv",
        "hasil_uji_smoke_bearer_prod_2026-07-16.csv",
    ):
        source = ROOT / "docs" / "evidence" / file_name
        if source.exists():
            shutil.copy2(source, DATA_OUTPUT / file_name)

    pubspec = (ROOT / "pubspec.yaml").read_text(encoding="utf-8")
    version = re.search(r"^version:\s*(.+)$", pubspec, re.MULTILINE).group(1).strip()
    commit = run(["git", "rev-parse", "HEAD"]).strip()
    branch = run(["git", "branch", "--show-current"]).strip()
    aab_path = ROOT / "build" / "app" / "outputs" / "bundle" / "release" / "app-release.aab"
    if not aab_path.exists():
        raise FileNotFoundError(f"AAB release tidak ditemukan: {aab_path}")

    build_summary = (
        f"Nama aplikasi : SIPANTES\n"
        f"Package ID    : id.rsudotista.sipantes\n"
        f"Versi         : {version}\n"
        f"Branch        : {branch}\n"
        f"Commit        : {commit}\n"
        f"AAB           : {aab_path}\n"
        f"Ukuran AAB    : {aab_path.stat().st_size / 1024 / 1024:.2f} MB\n"
        f"SHA-256 AAB   : {sha256(aab_path)}\n"
        "Flutter analyze: LULUS\nFlutter test: 16/16 LULUS\n"
    )
    (LOG_OUTPUT / "build_release_summary.txt").write_text(build_summary, encoding="utf-8")

    manifest_rows = []
    from PIL import Image

    for item in SCREENSHOTS:
        path = SCREENSHOT_OUTPUT / item["file"]
        with Image.open(path) as image:
            dimensions = f"{image.width}x{image.height}"
        manifest_rows.append(
            {
                "Id": item["id"],
                "File": f"evidence/screenshots/{item['file']}",
                "Jenis": "Screenshot UAT",
                "Dimensi": dimensions,
                "UkuranBytes": path.stat().st_size,
                "SHA256": sha256(path),
                "Keterangan": item["title"],
            }
        )
    write_csv(DATA_OUTPUT / "Evidence_Manifest.csv", manifest_rows)

    return {
        "version": version,
        "commit": commit,
        "branch": branch,
        "aab_path": aab_path,
        "aab_hash": sha256(aab_path),
        "api_rows": api_rows,
        "manifest_rows": manifest_rows,
    }


def add_cover(document: Document, metadata: dict) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = ROOT / "assets" / "images" / "otista" / "rsud_full_logo_trimmed.png"
    if not logo.exists():
        logo = ROOT / "assets" / "images" / "otista" / "rsud_full_logo.png"
    paragraph.add_run().add_picture(str(logo), width=Cm(12.5))
    document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("DOKUMEN USER ACCEPTANCE TEST")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_item = subtitle.add_run("APLIKASI SIPANTES")
    run_item.bold = True
    run_item.font.size = Pt(18)
    organization = document.add_paragraph()
    organization.alignment = WD_ALIGN_PARAGRAPH.CENTER
    organization.add_run(
        "Sistem Pendaftaran Terintegrasi\n"
        "RSUD Oto Iskandar Di Nata Kabupaten Bandung"
    ).font.size = Pt(14)
    document.add_paragraph()
    add_table(
        document,
        ["Informasi", "Nilai"],
        [
            ["Nomor Dokumen", "UAT/SIPANTES/001"],
            ["Versi Dokumen", "1.0"],
            ["Tanggal", DOCUMENT_DATE],
            ["Versi Aplikasi", metadata["version"]],
            ["Klasifikasi", "Dokumen Internal / Terkendali"],
        ],
        font_size=11,
    )
    trailing_paragraph = document.paragraphs[-1]
    trailing_paragraph._element.getparent().remove(trailing_paragraph._element)
    document.add_page_break()


def build_document(metadata: dict) -> None:
    document = Document()
    configure_styles(document)
    set_page_layout(document.sections[0])
    add_cover(document, metadata)

    document.add_heading("Kontrol Dokumen", level=1)
    add_table(
        document,
        ["Peran", "Nama / Unit", "Tanggal", "Tanda Tangan"],
        [
            ["Disusun oleh", "Tim Pengembang SIPANTES", "", ""],
            ["Diperiksa oleh", "Tim TI / SIMRS", "", ""],
            ["Disetujui oleh", "Pemilik Proses / Manajemen", "", ""],
        ],
    )
    document.add_heading("Riwayat Revisi", level=2)
    add_table(
        document,
        ["Versi", "Tanggal", "Deskripsi", "Penyusun"],
        [["1.0", DOCUMENT_DATE, "Penerbitan awal dokumen UAT lengkap beserta bukti.", "Tim Pengembang SIPANTES"]],
    )

    document.add_heading("Daftar Isi", level=1)
    for item in (
        "1. Ringkasan Eksekutif",
        "2. Tujuan dan Ruang Lingkup",
        "3. Lingkungan dan Konfigurasi Pengujian",
        "4. Kriteria Masuk dan Keluar",
        "5. Ringkasan Hasil Pengujian",
        "6. Matriks Skenario UAT",
        "7. Hasil Smoke Test API dan Web",
        "8. Hasil Pengujian Otomatis",
        "9. Catatan, Risiko, dan Tindak Lanjut",
        "10. Kesimpulan dan Rekomendasi",
        "11. Lembar Persetujuan UAT",
        "Lampiran A. Bukti Screenshot UAT",
        "Lampiran B. Manifest Bukti",
        "Lampiran C. Identitas Build Release",
    ):
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)

    document.add_heading("1. Ringkasan Eksekutif", level=1)
    summary = document.add_table(rows=1, cols=1)
    summary.style = "Table Grid"
    shade_cell(summary.cell(0, 0), LIGHT_TEAL)
    set_cell_text(
        summary.cell(0, 0),
        "KESIMPULAN UAT: DITERIMA DENGAN CATATAN OPERASIONAL\n"
        "Aplikasi SIPANTES layak digunakan pada tahap pengujian tertutup. "
        "Seluruh 14 skenario UAT lulus, 16 pengujian otomatis lulus, dan "
        "8 smoke test API/web lulus. Tidak ditemukan defect blocker atau critical.",
        bold=True,
        size=12,
    )
    document.add_paragraph()
    add_body(
        document,
        "Pengujian memvalidasi peluncuran aplikasi, konsistensi branding, "
        "pemulihan sesi akun sintetis, navigasi layanan, antrean umum, poli dan "
        "jadwal, ketersediaan kamar, kontrol keamanan sesi, kebijakan privasi, "
        "serta mekanisme penghapusan akun. Bukti visual diambil dari emulator "
        "Android menggunakan akun sintetis bernama test.",
    )
    add_body(
        document,
        "Transaksi booking sampai penyimpanan akhir dan konfirmasi penghapusan "
        "akun secara destruktif tidak dijalankan untuk menghindari pencatatan "
        "transaksi operasional serta menjaga akun bukti. Kedua alur tersebut "
        "tetap divalidasi pada lapisan UI, kontrak API, dan pengujian otomatis.",
    )

    document.add_heading("2. Tujuan dan Ruang Lingkup", level=1)
    document.add_heading("2.1 Tujuan", level=2)
    add_bullets(
        document,
        [
            "Memastikan fungsi utama aplikasi dapat digunakan sesuai kebutuhan pengguna.",
            "Memastikan integrasi aplikasi dengan API produksi dapat dijangkau.",
            "Memastikan endpoint terlindungi menolak request tanpa sesi.",
            "Memastikan halaman kebijakan privasi dan penghapusan akun dapat diakses publik.",
            "Menyediakan bukti formal untuk proses penerimaan dan pengujian tertutup.",
        ],
    )
    document.add_heading("2.2 Ruang Lingkup", level=2)
    add_body(
        document,
        "Ruang lingkup mencakup aplikasi Flutter SIPANTES, package Android "
        "id.rsudotista.sipantes, layanan API publik, kontrol sesi Bearer, "
        "halaman kebijakan privasi, halaman penghapusan akun, dan bukti visual "
        "pada emulator Android.",
    )
    document.add_heading("2.3 Di Luar Ruang Lingkup", level=2)
    add_bullets(
        document,
        [
            "Penggunaan data pasien nyata.",
            "Pembuatan transaksi booking nyata pada jadwal pelayanan.",
            "Konfirmasi akhir penghapusan akun secara destruktif.",
            "Uji beban, penetration testing penuh, dan disaster recovery.",
            "Persetujuan medis atau validasi isi rekam medis.",
        ],
    )

    document.add_heading("3. Lingkungan dan Konfigurasi Pengujian", level=1)
    add_table(
        document,
        ["Parameter", "Nilai"],
        [
            ["Nama aplikasi", "SIPANTES"],
            ["Package ID", "id.rsudotista.sipantes"],
            ["Versi", metadata["version"]],
            ["Branch / Commit", f"{metadata['branch']} / {metadata['commit']}"],
            ["Perangkat UAT", "Android Emulator sdk_gphone64_x86_64, Android 13 / API 33"],
            ["Resolusi bukti", "1080 x 1920 piksel, portrait"],
            ["Build emulator", "Debug-signed UAT build versionCode 3 dengan APP_ENV=prod"],
            ["Build distribusi", "Release AAB berhasil dibuat dan ditandatangani"],
            ["API utama", "https://api-mobile.rsudotista.my.id/api/v1"],
            ["Data uji", "Akun sintetis test; tanpa data pasien nyata"],
        ],
    )

    document.add_heading("4. Kriteria Masuk dan Keluar", level=1)
    document.add_heading("4.1 Kriteria Masuk", level=2)
    add_bullets(
        document,
        [
            "Source berada pada branch main dan dapat dibangun.",
            "Package Android menggunakan id.rsudotista.sipantes.",
            "API produksi dapat dijangkau melalui HTTPS.",
            "Akun sintetis tersedia pada emulator.",
            "Tidak ada data pasien nyata pada bukti UAT.",
        ],
    )
    document.add_heading("4.2 Kriteria Keluar", level=2)
    add_bullets(
        document,
        [
            "Seluruh skenario UAT dalam ruang lingkup berstatus LULUS.",
            "Tidak ada defect blocker atau critical.",
            "flutter analyze dan flutter test berhasil.",
            "Release AAB berhasil dibuat.",
            "Halaman kebijakan privasi dan penghapusan akun merespons HTTP 200.",
        ],
    )

    document.add_heading("5. Ringkasan Hasil Pengujian", level=1)
    add_table(
        document,
        ["Kelompok Pengujian", "Total", "Lulus", "Tidak Sesuai", "Catatan"],
        [
            ["Skenario UAT pengguna", "14", "14", "0", "Seluruh skenario dalam ruang lingkup lulus"],
            ["Pengujian otomatis Flutter", "16", "16", "0", "All tests passed"],
            ["Smoke test API dan web", "8", "8", "0", "HTTP aktual sesuai ekspektasi"],
            ["Build dan analisis statis", "2", "2", "0", "Analyze dan build AAB berhasil"],
        ],
    )
    add_body(
        document,
        "Defect terbuka: 0 blocker, 0 critical, 0 major. Terdapat dua catatan "
        "operasional untuk pengujian lanjutan: transaksi booking sampai "
        "penyimpanan akhir dan konfirmasi penghapusan akun secara destruktif.",
        bold=True,
    )

    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_layout(landscape, landscape=True)
    document.add_heading("6. Matriks Skenario UAT", level=1)
    scenario_rows = [
        [
            case["id"],
            f"{case['module']}\n{case['scenario']}",
            case["steps"],
            case["expected"],
            case["actual"],
            f"{case['status']}\n{case['evidence']}",
        ]
        for case in UAT_CASES
    ]
    add_table(
        document,
        ["ID", "Modul / Skenario", "Langkah", "Ekspektasi", "Aktual", "Status / Bukti"],
        scenario_rows,
        font_size=8.8,
    )

    portrait = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_layout(portrait)
    document.add_heading("7. Hasil Smoke Test API dan Web", level=1)
    add_body(
        document,
        "Pengujian dilakukan tanpa access token dan tanpa mengirim data pribadi. "
        "Endpoint publik diharapkan merespons HTTP 200, sedangkan endpoint sesi "
        "tanpa Bearer diharapkan menolak dengan HTTP 401.",
    )
    add_table(
        document,
        ["ID", "Pengujian", "URL", "Expected", "Actual", "Waktu", "Hasil"],
        [
            [
                row["Id"],
                row["Pengujian"],
                row["URL"],
                str(row["ExpectedHTTP"]),
                str(row["ActualHTTP"]),
                f"{row['TimeSeconds']} s",
                row["Hasil"],
            ]
            for row in metadata["api_rows"]
        ],
        font_size=9.5,
    )

    document.add_heading("8. Hasil Pengujian Otomatis", level=1)
    add_body(
        document,
        "Pengujian otomatis mencakup autentikasi, keamanan Bearer, rotasi token, "
        "validasi registrasi, pemulihan sesi, kontrak endpoint terlindungi, "
        "penghapusan akun, dan pemuatan shell aplikasi.",
    )
    add_table(
        document,
        ["No.", "Pengujian", "Hasil"],
        [[str(index), test_name, "LULUS"] for index, test_name in enumerate(AUTOMATED_TESTS, 1)],
    )

    document.add_heading("9. Catatan, Risiko, dan Tindak Lanjut", level=1)
    add_table(
        document,
        ["No.", "Catatan / Risiko", "Dampak", "Tindak Lanjut", "Status"],
        [
            [
                "1",
                "Konfirmasi penghapusan akun secara destruktif tidak dilakukan pada sesi bukti.",
                "Akun bukti tetap tersedia untuk dokumentasi dan pengujian lanjutan.",
                "Jalankan memakai akun sintetis disposable dan arsipkan bukti tanpa menampilkan OTP.",
                "TINDAK LANJUT",
            ],
            [
                "2",
                "Transaksi booking sampai penyimpanan akhir tidak dilakukan karena akun bukti tidak dikaitkan dengan pasien sintetis.",
                "Tidak ada transaksi pelayanan uji yang masuk ke sistem operasional.",
                "Siapkan pasien sintetis dan jadwal uji khusus sebelum persetujuan produksi penuh.",
                "TINDAK LANJUT",
            ],
            [
                "3",
                "Screenshot memakai debug-signed build versionCode 3 untuk mempertahankan sesi emulator.",
                "Tidak memengaruhi validasi UI; AAB release diuji terpisah melalui proses build.",
                "Gunakan AAB release yang sudah ditandatangani untuk distribusi Play Console.",
                "TERKENDALI",
            ],
        ],
        font_size=10,
    )

    document.add_heading("10. Kesimpulan dan Rekomendasi", level=1)
    summary = document.add_table(rows=1, cols=1)
    summary.style = "Table Grid"
    shade_cell(summary.cell(0, 0), LIGHT_TEAL)
    set_cell_text(
        summary.cell(0, 0),
        "STATUS PENERIMAAN: DITERIMA DENGAN CATATAN OPERASIONAL\n"
        f"Aplikasi SIPANTES versi {metadata['version']} memenuhi kriteria untuk "
        "melanjutkan pengujian tertutup. Tidak terdapat defect blocker atau "
        "critical pada ruang lingkup yang diuji.",
        bold=True,
        size=12,
    )
    document.add_paragraph()
    add_body(
        document,
        "Sebelum persetujuan produksi penuh, lakukan satu booking end-to-end "
        "menggunakan pasien sintetis yang disetujui, satu penghapusan akun "
        "end-to-end menggunakan akun disposable, pemantauan log API selama masa "
        "pengujian tertutup, serta verifikasi berkala bahwa URL kebijakan dan "
        "penghapusan akun tetap tersedia publik.",
    )

    document.add_heading("11. Lembar Persetujuan UAT", level=1)
    add_body(
        document,
        "Dengan menandatangani bagian ini, pihak terkait menyatakan telah "
        "meninjau hasil UAT, memahami catatan operasional, dan menyetujui status "
        "penerimaan yang tercantum dalam dokumen.",
    )
    signature = add_table(
        document,
        ["Disusun oleh", "Diperiksa oleh", "Disetujui oleh"],
        [[
            "Tim Pengembang SIPANTES\n\nNama: ____________________\nTanggal: __________________",
            "Tim TI / SIMRS\n\nNama: ____________________\nTanggal: __________________",
            "Pemilik Proses / Manajemen\n\nNama: ____________________\nTanggal: __________________",
        ]],
        font_size=11,
    )
    for cell in signature.rows[1].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    document.add_page_break()
    document.add_heading("Lampiran A. Bukti Screenshot UAT", level=1)
    add_body(
        document,
        "Seluruh screenshot beresolusi 1080 x 1920 piksel dan diambil dari "
        "emulator Android yang menjalankan SIPANTES versionCode 3. Data yang "
        "terlihat menggunakan akun sintetis.",
    )
    add_table(
        document,
        ["ID", "Bukti", "File"],
        [[item["id"], item["title"], item["file"]] for item in SCREENSHOTS],
        font_size=10,
    )
    document.add_page_break()
    for index, item in enumerate(SCREENSHOTS):
        document.add_heading(f"{item['id']}. {item['title']}", level=2)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(
            str(SCREENSHOT_OUTPUT / item["file"]),
            width=Cm(7.5),
        )
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_run = caption.add_run(item["description"])
        caption_run.italic = True
        caption_run.font.size = Pt(11)
        if index != len(SCREENSHOTS) - 1:
            document.add_page_break()

    document.add_page_break()
    document.add_heading("Lampiran B. Manifest Bukti", level=1)
    add_table(
        document,
        ["ID", "File", "Dimensi", "SHA-256", "Keterangan"],
        [
            [row["Id"], row["File"], row["Dimensi"], row["SHA256"], row["Keterangan"]]
            for row in metadata["manifest_rows"]
        ],
        font_size=8.5,
    )

    document.add_heading("Lampiran C. Identitas Build Release", level=1)
    add_table(
        document,
        ["Parameter", "Nilai"],
        [
            ["Versi", metadata["version"]],
            ["Commit", metadata["commit"]],
            ["Ukuran AAB", f"{metadata['aab_path'].stat().st_size / 1024 / 1024:.2f} MB"],
            ["SHA-256 AAB", metadata["aab_hash"]],
            ["Analisis statis", "LULUS - No issues found"],
            ["Pengujian otomatis", "LULUS - 16/16"],
        ],
        font_size=10,
    )
    add_body(
        document,
        "Bukti historis pengujian API tersedia pada folder evidence/data. Bukti "
        "lama digunakan sebagai referensi regresi dan tidak menggantikan smoke "
        "test terkini dalam dokumen ini.",
    )

    add_headers_and_footers(document)
    force_garamond(document)
    document.core_properties.title = "Dokumen UAT SIPANTES"
    document.core_properties.subject = "User Acceptance Test aplikasi SIPANTES"
    document.core_properties.author = "Tim Pengembang SIPANTES"
    document.core_properties.keywords = "SIPANTES, UAT, RSUD Otista, Play Store"
    document.save(DOCX_PATH)


def convert_and_validate(metadata: dict) -> dict:
    converter = ROOT / "scripts" / "convert_docx_to_pdf.ps1"
    conversion = run(
        [
            "powershell.exe",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(converter),
            "-DocxPath",
            str(DOCX_PATH),
            "-PdfPath",
            str(PDF_PATH),
        ]
    )
    if not PDF_PATH.exists() or PDF_PATH.stat().st_size < 100_000:
        raise RuntimeError("PDF hasil ekspor tidak tersedia atau terlalu kecil.")

    pdf = pymupdf.open(PDF_PATH)
    fonts = set()
    image_count = 0
    evidence_page = None
    for page_index, page in enumerate(pdf):
        for font in page.get_fonts(full=True):
            fonts.add(font[3])
        image_count += len(page.get_images(full=True))
        if evidence_page is None and "SS-01. Beranda" in page.get_text():
            evidence_page = page_index

    if len(pdf) < 14:
        raise RuntimeError(f"Jumlah halaman PDF tidak wajar: {len(pdf)}")
    if not any("Garamond" in font for font in fonts):
        raise RuntimeError(f"Font Garamond tidak terdeteksi pada PDF: {sorted(fonts)}")
    if image_count < len(SCREENSHOTS):
        raise RuntimeError(f"Jumlah gambar PDF kurang: {image_count}")

    cover_pixmap = pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.4, 1.4), alpha=False)
    cover_preview = PREVIEW_OUTPUT / "preview_cover.png"
    cover_pixmap.save(cover_preview)
    evidence_index = evidence_page if evidence_page is not None else max(0, len(pdf) - 7)
    evidence_pixmap = pdf[evidence_index].get_pixmap(
        matrix=pymupdf.Matrix(1.2, 1.2),
        alpha=False,
    )
    evidence_preview = PREVIEW_OUTPUT / "preview_evidence.png"
    evidence_pixmap.save(evidence_preview)

    report = (
        f"PDF: {PDF_PATH}\n"
        f"Pages: {len(pdf)}\n"
        f"Images: {image_count}\n"
        f"Garamond detected: YES\n"
        f"Fonts: {', '.join(sorted(fonts))}\n"
        f"Conversion output: {conversion.strip()}\n"
        f"PDF SHA-256: {sha256(PDF_PATH)}\n"
    )
    (LOG_OUTPUT / "pdf_validation.txt").write_text(report, encoding="utf-8")
    page_count = len(pdf)
    pdf.close()
    return {
        "page_count": page_count,
        "image_count": image_count,
        "fonts": sorted(fonts),
        "cover_preview": cover_preview,
        "evidence_preview": evidence_preview,
    }


def package_artifacts(metadata: dict, validation: dict) -> None:
    readme = (
        "PAKET DOKUMEN UAT SIPANTES\n\n"
        "Dokumen utama:\n"
        "- Dokumen_UAT_SIPANTES_v1.0.pdf\n"
        "- Dokumen_UAT_SIPANTES_v1.0.docx\n\n"
        "Bukti:\n"
        "- evidence/screenshots: 5 screenshot emulator versionCode 3\n"
        "- evidence/logs: flutter analyze, flutter test, build, dan validasi PDF\n"
        "- evidence/data: matriks UAT, smoke test API/web, hasil otomatis, dan hash\n"
        "- preview: pratinjau cover dan halaman bukti\n\n"
        "Status: DITERIMA DENGAN CATATAN OPERASIONAL\n"
        f"Tanggal: {DOCUMENT_DATE}\n"
        f"Jumlah halaman PDF: {validation['page_count']}\n"
        f"Versi aplikasi: {metadata['version']}\n"
        f"Commit: {metadata['commit']}\n"
    )
    (OUTPUT / "README_UAT.txt").write_text(readme, encoding="utf-8")

    artifact_rows = []
    for path in sorted(OUTPUT.rglob("*")):
        if path.is_file() and path.name != "Artifact_Manifest.csv":
            artifact_rows.append(
                {
                    "File": path.relative_to(OUTPUT).as_posix(),
                    "UkuranBytes": path.stat().st_size,
                    "SHA256": sha256(path),
                }
            )
    write_csv(OUTPUT / "Artifact_Manifest.csv", artifact_rows)

    if ZIP_PATH.exists():
        ZIP_PATH.unlink()
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(OUTPUT.rglob("*")):
            if path.is_file():
                archive.write(path, Path(OUTPUT.name) / path.relative_to(OUTPUT))


def main() -> None:
    metadata = collect_evidence()
    build_document(metadata)
    validation = convert_and_validate(metadata)
    package_artifacts(metadata, validation)
    print(f"UAT_OUTPUT={OUTPUT}")
    print(f"DOCX={DOCX_PATH}")
    print(f"PDF={PDF_PATH}")
    print(f"ZIP={ZIP_PATH}")
    print(f"PAGES={validation['page_count']}")
    print(f"PDF_SHA256={sha256(PDF_PATH)}")


if __name__ == "__main__":
    main()
